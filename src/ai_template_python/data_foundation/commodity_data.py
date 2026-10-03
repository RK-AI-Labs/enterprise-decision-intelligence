"""Import the World Bank Pink Sheet monthly commodity workbook."""

import hashlib
import re
from collections.abc import Iterator
from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

PINK_SHEET_URL = "https://www.worldbank.org/en/research/commodity-markets"
_MONTH_KEY = re.compile(r"^(?P<year>\d{4})M(?P<month>\d{2})$")


@dataclass(frozen=True, slots=True)
class CommodityObservation:
    source_name: str
    series_id: str
    country_code: str
    observed_on: date
    value: Decimal
    unit: str
    currency: str | None
    source_url: str
    retrieved_at: datetime
    source_version: str


def _series_id(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", name.casefold()).strip("-")


def iter_pink_sheet_monthly(path: str | Path) -> Iterator[CommodityObservation]:
    """Yield observed numeric cells from the Pink Sheet ``Monthly Prices`` sheet."""
    workbook_path = Path(path)
    source_version = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    retrieved_at = datetime.now(UTC)
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if "Monthly Prices" not in workbook.sheetnames:
            raise ValueError("World Bank workbook has no 'Monthly Prices' sheet")
        sheet = workbook["Monthly Prices"]
        rows = sheet.iter_rows(min_row=5, max_row=6, values_only=True)
        names = next(rows)
        units = next(rows)
        if not names or not units or names[0] is not None:
            raise ValueError("Unexpected World Bank Pink Sheet column header layout")

        series = [
            (index, str(name).strip(), str(units[index] or "unspecified"))
            for index, name in enumerate(names)
            if index > 0 and name is not None and str(name).strip()
        ]
        for row in sheet.iter_rows(min_row=7, values_only=True):
            if not row or not isinstance(row[0], str):
                continue
            month_match = _MONTH_KEY.fullmatch(row[0].strip())
            if month_match is None:
                continue
            observed_on = date(
                int(month_match.group("year")),
                int(month_match.group("month")),
                1,
            )
            for column, name, unit in series:
                value = row[column] if column < len(row) else None
                if not isinstance(value, (int, float, Decimal)):
                    continue
                yield CommodityObservation(
                    source_name="world_bank_pink_sheet",
                    series_id=_series_id(name),
                    country_code="WLD",
                    observed_on=observed_on,
                    value=Decimal(str(value)),
                    unit=unit,
                    currency="USD" if "$" in unit else None,
                    source_url=PINK_SHEET_URL,
                    retrieved_at=retrieved_at,
                    source_version=source_version,
                )
    finally:
        workbook.close()
